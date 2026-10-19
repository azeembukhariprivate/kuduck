"""KUDUCK document helper.

Converts PDF files to Word (.docx) and Excel (.xlsx). The KUDUCK app runs this
as a small background program and reads one line of JSON from its output.

Usage:
    kuduck_helper pdf2word <input.pdf> <output.docx> [password]
    kuduck_helper pdf2excel <input.pdf> <output.xlsx> [password]

Exit codes: 0 ok, 2 bad arguments, 3 scanned PDF (no text), 4 wrong password,
1 any other error.
"""
import io
import json
import re
import sys
import traceback

import pdfplumber
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, Inches
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def out(ok, message, code=0, **extra):
    print(json.dumps({"ok": ok, "message": message, **extra}), flush=True)
    sys.exit(code)


def open_pdf(path, password):
    try:
        return pdfplumber.open(path, password=password or "")
    except Exception as exc:  # pdfplumber wraps pdfminer's password errors
        names = " ".join(type(e).__name__ for e in (exc, exc.__cause__, exc.__context__) if e)
        names += " " + " ".join(type(a).__name__ for a in getattr(exc, "args", ()))
        text = (names + " " + str(exc)).lower()
        if "password" in text or "encrypt" in text:
            out(False, "This PDF is password protected. Unlock it first or enter the password.", 4)
        raise


def has_text(pdf):
    for page in pdf.pages[: min(len(pdf.pages), 5)]:
        if page.chars:
            return True
    return False


# ---------------------------------------------------------------- PDF -> Word

def _inside(bbox, box):
    x0, top, x1, bottom = bbox
    return x0 >= box[0] - 1 and x1 <= box[2] + 1 and top >= box[1] - 1 and bottom <= box[3] + 1


def _lines(page, skip_boxes):
    words = page.extract_words(keep_blank_chars=False, use_text_flow=True,
                               extra_attrs=["size", "fontname"])
    words = [w for w in words if not any(_inside((w["x0"], w["top"], w["x1"], w["bottom"]), b)
                                         for b in skip_boxes)]
    lines = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if lines and abs(lines[-1]["top"] - w["top"]) < max(2.0, w["size"] * 0.3):
            ln = lines[-1]
            ln["words"].append(w)
            ln["bottom"] = max(ln["bottom"], w["bottom"])
        else:
            lines.append({"top": w["top"], "bottom": w["bottom"], "words": [w]})
    for ln in lines:
        ln["words"].sort(key=lambda w: w["x0"])
        ln["text"] = " ".join(w["text"] for w in ln["words"])
        ln["size"] = max(w["size"] for w in ln["words"])
        ln["bold"] = all("bold" in w["fontname"].lower() or "black" in w["fontname"].lower()
                         for w in ln["words"])
        ln["x0"] = ln["words"][0]["x0"]
    return lines


def _body_size(pdf):
    sizes = {}
    for page in pdf.pages[: min(len(pdf.pages), 10)]:
        for c in page.chars:
            s = round(c["size"])
            sizes[s] = sizes.get(s, 0) + 1
    return max(sizes, key=sizes.get) if sizes else 11


def pdf_to_word(src, dst, password=None):
    with open_pdf(src, password) as pdf:
        if not has_text(pdf):
            out(False, "This PDF looks scanned (no text). Use Scan → Text instead.", 3)
        body = _body_size(pdf)
        doc = Document()
        style = doc.styles["Normal"]
        style.font.size = Pt(max(8, min(14, body)))

        for pno, page in enumerate(pdf.pages):
            if pno > 0:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            tables = page.find_tables()
            table_boxes = [t.bbox for t in tables]
            blocks = []
            for t in tables:
                rows = t.extract()
                if rows:
                    blocks.append(("table", t.bbox[1], rows))
            for im in page.images:
                bbox = (max(im["x0"], 0), max(im["top"], 0),
                        min(im["x1"], page.width), min(im["bottom"], page.height))
                if bbox[2] - bbox[0] > 20 and bbox[3] - bbox[1] > 20:
                    blocks.append(("image", bbox[1], bbox))
            for ln in _lines(page, table_boxes):
                blocks.append(("line", ln["top"], ln))
            blocks.sort(key=lambda b: b[1])

            para = None
            last = None
            for kind, _, data in blocks:
                if kind == "table":
                    ncols = max(len(r) for r in data)
                    tbl = doc.add_table(rows=len(data), cols=ncols)
                    tbl.style = "Table Grid"
                    for r, row in enumerate(data):
                        for c in range(ncols):
                            val = row[c] if c < len(row) and row[c] is not None else ""
                            tbl.cell(r, c).text = str(val)
                    para, last = None, None
                elif kind == "image":
                    try:
                        pic = page.crop(data).to_image(resolution=150).original
                        buf = io.BytesIO()
                        pic.save(buf, format="PNG")
                        buf.seek(0)
                        width_in = min(6.5, (data[2] - data[0]) / 72.0)
                        doc.add_picture(buf, width=Inches(width_in))
                    except Exception:
                        pass
                    para, last = None, None
                else:
                    ln = data
                    heading = ln["size"] >= body + 2.5
                    gap = (ln["top"] - last["bottom"]) if last else 999
                    same_para = (para is not None and not heading and not last.get("heading")
                                 and gap < ln["size"] * 0.9
                                 and abs(ln["size"] - last["size"]) < 1)
                    if same_para:
                        para.add_run(" " + ln["text"]).bold = ln["bold"] or None
                    else:
                        para = doc.add_paragraph()
                        run = para.add_run(ln["text"])
                        run.font.size = Pt(round(ln["size"]))
                        if heading or ln["bold"]:
                            run.bold = True
                    ln["heading"] = heading
                    last = ln
        doc.save(dst)
        return len(pdf.pages)


# ---------------------------------------------------------------- PDF -> Excel

_NUM = re.compile(r"^-?[\d,]*\.?\d+$")


def _cell_value(text):
    if text is None:
        return None
    t = str(text).strip()
    if _NUM.match(t.replace(" ", "")):
        try:
            v = float(t.replace(",", "").replace(" ", ""))
            return int(v) if v.is_integer() and "." not in t else v
        except ValueError:
            pass
    return t


def pdf_to_excel(src, dst, password=None):
    with open_pdf(src, password) as pdf:
        if not has_text(pdf):
            out(False, "This PDF looks scanned (no text). Use Scan → Text instead.", 3)
        wb = Workbook()
        wb.remove(wb.active)
        found = 0
        for pno, page in enumerate(pdf.pages, start=1):
            tables = [t for t in page.extract_tables() if t and len(t) > 1]
            if not tables:
                tables = [t for t in page.extract_tables(
                    {"vertical_strategy": "text", "horizontal_strategy": "text"}) if t and len(t) > 1]
            ws = wb.create_sheet(f"Page {pno}")
            row = 1
            widths = {}
            if tables:
                for t in tables:
                    found += 1
                    for r in t:
                        for c, val in enumerate(r, start=1):
                            v = _cell_value(val)
                            ws.cell(row=row, column=c, value=v)
                            widths[c] = max(widths.get(c, 0), len(str(v or "")))
                        row += 1
                    row += 1
            else:
                for line in (page.extract_text() or "").splitlines():
                    ws.cell(row=row, column=1, value=line)
                    widths[1] = max(widths.get(1, 0), len(line))
                    row += 1
            for c, w in widths.items():
                ws.column_dimensions[get_column_letter(c)].width = min(60, max(8, w + 2))
        if not wb.sheetnames:
            wb.create_sheet("Page 1")
        wb.save(dst)
        return found


def main(argv):
    if len(argv) < 4 or argv[1] not in ("pdf2word", "pdf2excel"):
        out(False, "usage: kuduck_helper pdf2word|pdf2excel <in.pdf> <out> [password]", 2)
    cmd, src, dst = argv[1], argv[2], argv[3]
    password = argv[4] if len(argv) > 4 else None
    try:
        if cmd == "pdf2word":
            pages = pdf_to_word(src, dst, password)
            out(True, f"Converted {pages} page(s) to Word.", output=dst)
        else:
            tables = pdf_to_excel(src, dst, password)
            msg = f"Found {tables} table(s)." if tables else "No tables found; text copied into the sheets."
            out(True, msg, output=dst)
    except SystemExit:
        raise
    except PermissionError:
        out(False, "No permission to save there. Choose another folder, or run KUDUCK as administrator.", 1)
    except Exception as exc:
        traceback.print_exc(file=sys.stderr)
        out(False, f"Conversion failed: {exc}", 1)


if __name__ == "__main__":
    main(sys.argv)
